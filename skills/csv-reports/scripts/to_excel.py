#!/usr/bin/env python3
"""Turn a CSV into a workbook somebody can open and send on.

A bold frozen header, columns wide enough for what is in them, numbers right
aligned, and a totals row under any column that is entirely numeric.

Needs openpyxl, which is not part of Python. If it is missing this exits with a
message naming it: install it (references/setup.md) and run this again rather
than rewriting the script to avoid it.
"""

import csv
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
except ModuleNotFoundError as missing:
    # Named, with the fix, because this is the message somebody has to act on.
    print(
        f"this script needs the {missing.name} library and it is not installed.\n"
        f"install it with:  python3 -m pip install {missing.name}\n"
        f"if that answers 'externally-managed-environment', see references/setup.md",
        file=sys.stderr,
    )
    raise SystemExit(3)


def looks_numeric(value):
    """Whether a cell is a number. Blank is not a number and not a failure."""
    if value is None or value.strip() == "":
        return None
    try:
        return float(value.replace(",", ""))
    except ValueError:
        return False


def main(argv):
    if len(argv) < 2:
        print("usage: to_excel.py <input.csv> [output.xlsx]", file=sys.stderr)
        return 2

    source = Path(argv[1])
    if not source.is_file():
        print(f"there is no file at {source}", file=sys.stderr)
        return 1
    target = Path(argv[2]) if len(argv) > 2 else source.with_suffix(".xlsx")

    with source.open(newline="", encoding="utf-8-sig") as handle:
        rows = list(csv.reader(handle))
    if not rows:
        print(f"{source} is empty", file=sys.stderr)
        return 1

    header, body = rows[0], rows[1:]
    book = Workbook()
    sheet = book.active
    sheet.title = source.stem[:31] or "Sheet1"

    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    # So the header stays put when somebody scrolls a thousand rows.
    sheet.freeze_panes = "A2"

    # Which columns are numbers all the way down. A blank does not disqualify a
    # column: a missing price is missing, not text.
    numeric = [True] * len(header)
    for row in body:
        for index, value in enumerate(row[: len(header)]):
            if looks_numeric(value) is False:
                numeric[index] = False

    for row in body:
        padded = list(row[: len(header)]) + [""] * (len(header) - len(row))
        written = []
        for index, value in enumerate(padded):
            number = looks_numeric(value) if numeric[index] else False
            written.append(number if isinstance(number, float) else value)
        sheet.append(written)

    # A totals row, only under the columns that have something to total.
    if body and any(numeric):
        totals = []
        for index, is_number in enumerate(numeric):
            if index == 0:
                totals.append("Total")
            elif is_number:
                letter = get_column_letter(index + 1)
                totals.append(f"=SUM({letter}2:{letter}{len(body) + 1})")
            else:
                totals.append("")
        sheet.append(totals)
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)

    for index, name in enumerate(header, start=1):
        longest = len(str(name))
        for row in body:
            if index <= len(row):
                longest = max(longest, len(str(row[index - 1])))
        sheet.column_dimensions[get_column_letter(index)].width = min(longest + 2, 60)
        if numeric[index - 1]:
            for cell in sheet[get_column_letter(index)]:
                cell.alignment = Alignment(horizontal="right")

    book.save(target)
    print(f"wrote {target} — sheet {sheet.title!r}, {len(body)} rows, {len(header)} columns")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
